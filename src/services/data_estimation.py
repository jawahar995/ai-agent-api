from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import math
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .estimator import estimate_text
from .node_lookup import fetch_data_record, save_estimate_result


@dataclass
class DataEstimateResult:
    data_id: str
    client_id: str | None
    source_type: str
    source_url: str
    page_count: int
    token_count: int
    chunk_count: int
    estimated_cost_usd: float
    estimated_seconds: float
    model: str


def _file_extension(source_url: str, data_type: str | None = None) -> str:
    if data_type:
        normalized = data_type.lower().strip()
        if normalized.startswith('.'):
            return normalized[1:]
        return normalized

    parsed = urlparse(source_url)
    ext = Path(parsed.path).suffix.lower().lstrip('.')
    return ext


async def _download_bytes(source_url: str) -> bytes:
    async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as client:
        response = await client.get(source_url)
        response.raise_for_status()
        return response.content


async def _download_text(source_url: str) -> str:
    async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as client:
        response = await client.get(source_url)
        response.raise_for_status()
        return response.text


async def _extract_text_and_pages(source_url: str, source_type: str | None, model: str | None) -> tuple[str, int]:
    ext = _file_extension(source_url, source_type)

    if ext in {"url", "html", "htm"}:
        html = await _download_text(source_url)
        soup = BeautifulSoup(html, "html.parser")
        title = soup.title.string.strip() if soup.title and soup.title.string else "Untitled"
        text = " ".join(item.strip() for item in soup.stripped_strings)
        page_count = max(1, math.ceil(len(text) / 3000))
        return f"# {title}\n\n{text}", page_count

    blob = await _download_bytes(source_url)

    if ext == "pdf":
        reader = PdfReader(BytesIO(blob))
        pages = len(reader.pages)
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        return text.strip(), max(1, pages)

    if ext in {"docx"}:
        document = Document(BytesIO(blob))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)
        page_count = max(1, math.ceil(len(text) / 3000))
        return text.strip(), page_count

    if ext in {"xlsx", "xlsm", "xltx", "xltm"}:
        workbook = load_workbook(BytesIO(blob), read_only=True, data_only=True)
        sheet_texts: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell) for cell in row if cell is not None]
                if row_values:
                    rows.append(" | ".join(row_values))
            if rows:
                sheet_texts.append(f"## {sheet.title}\n" + "\n".join(rows))
        text = "\n\n".join(sheet_texts)
        page_count = max(1, len(workbook.worksheets))
        return text.strip(), page_count

    if ext in {"csv", "txt", "md", "json", "xml"}:
        try:
            text = blob.decode("utf-8")
        except UnicodeDecodeError:
            text = blob.decode("utf-8", errors="ignore")
        page_count = max(1, math.ceil(len(text) / 3000))
        return text.strip(), page_count

    try:
        text = blob.decode("utf-8")
    except UnicodeDecodeError:
        text = blob.decode("utf-8", errors="ignore")
    page_count = max(1, math.ceil(len(text) / 3000))
    return text.strip(), page_count


async def estimate_from_data_id(data_id: str, model: str | None = None) -> DataEstimateResult:
    # 1. Fetch data record from Node (contains src_url, type, user_id)
    data = await fetch_data_record(data_id)
    source_url = data.get("src_url")
    if not source_url:
        raise FileNotFoundError(f"source url not found for data id: {data_id}")

    # 2. Download content and calculate estimation
    source_type = str(data.get("type") or "file")
    source_text, page_count = await _extract_text_and_pages(source_url, source_type, model)
    estimate = estimate_text(source_text, model=model)

    result = DataEstimateResult(
        data_id=data_id,
        client_id=str(data.get("user_id") or ""),
        source_type=source_type,
        source_url=source_url,
        page_count=page_count,
        token_count=estimate.token_count,
        chunk_count=estimate.chunk_count,
        estimated_cost_usd=estimate.estimated_cost_usd,
        estimated_seconds=estimate.estimated_seconds,
        model=estimate.model,
    )

    # 3. Persist estimation fields back to DB via Node
    await save_estimate_result(data_id, {
        "total_tokens": result.token_count,
        "total_chunks": result.chunk_count,
        "total_pages": result.page_count,
        "estimated_cost_usd": result.estimated_cost_usd,
        "estimated_seconds": result.estimated_seconds,
    })

    return result
